"""Virginia Department of Health, Beach Monitoring Program.

Findings (Oct 2026):
  * The public map is an embedded Tableau Public dashboard
    (workbook "2024BEACHMonitoringMap") -- awkward to scrape.
  * BUT the "Beach Data" page publishes per-year Excel workbooks:
        https://www.vdh.virginia.gov/waterborne-hazards-control/statistics/
    - Monitoring data: BeachMonitoringData_<stamp>.xlsx  (2018-2026)
    - Advisory data:   BeachActivityReport_<stamp>.xlsx  (2018-2026)
    The 2026 files were refreshed 2026-10-05. File names change with each
    refresh, so we scrape the page for the link labeled with the year.
  * Column names have not been inspected yet (binary download blocked in
    the research environment), so the parser locates the header row and
    maps columns by fuzzy name matching. Verify on first live run.
"""
from __future__ import annotations

import io
import re
from datetime import date
from typing import Optional

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from ..models import Advisory, CollectorResult, Indicator, Sample, Site
from .base import Collector, find_key, parse_date, parse_result, slug

PAGE = "https://www.vdh.virginia.gov/waterborne-hazards-control/statistics/"

NAME_KEYS = ("beach name", "beach", "site name", "station name", "location")
STATION_KEYS = ("station id", "station", "site id", "sample site", "beach id")
DATE_KEYS = ("sample date", "collection date", "date sampled", "date")
RESULT_KEYS = ("enterococci", "enterococcus", "entero", "result", "cfu")
LAT_KEYS = ("latitude", "lat")
LON_KEYS = ("longitude", "long", "lon")
COUNTY_KEYS = ("locality", "county", "city", "health district", "district")
ADV_START = ("advisory start", "start date", "date issued", "issued", "begin")
ADV_END = ("advisory end", "end date", "date lifted", "lifted", "rescinded")
ADV_TYPE = ("advisory type", "action type", "type", "action")
ADV_REASON = ("reason", "cause", "comment")


class VirginiaCollector(Collector):
    source_id = "va_vdh"
    state = "VA"
    agency = "Virginia Department of Health, Beach Monitoring Program"
    homepage_url = PAGE

    def __init__(self, year: Optional[int] = None, **kw):
        super().__init__(**kw)
        self.year = year or date.today().year

    def collect(self) -> CollectorResult:
        links = find_workbook_links(self.get(PAGE).text)
        res = CollectorResult()
        mon = links["monitoring"].get(self.year)
        if mon:
            res.extend(parse_monitoring_xlsx(self.get(mon).body, detail_url=PAGE))
        else:
            res.warnings.append(f"VA: no monitoring workbook for {self.year}")
        adv = links["advisory"].get(self.year)
        if adv:
            res.extend(parse_advisory_xlsx(self.get(adv).body))
        return res


def find_workbook_links(html: str) -> dict[str, dict[int, str]]:
    soup = BeautifulSoup(html, "lxml")
    out: dict[str, dict[int, str]] = {"monitoring": {}, "advisory": {}}
    for a in soup.find_all("a", href=True):
        href, label = a["href"], a.get_text(strip=True)
        if not href.lower().endswith((".xlsx", ".xls")) or not re.fullmatch(r"\d{4}", label):
            continue
        kind = ("monitoring" if "monitoring" in href.lower()
                else "advisory" if "activity" in href.lower() or "advisor" in href.lower()
                else None)
        if kind:
            out[kind][int(label)] = href
    return out


def _sheet_records(xlsx: bytes) -> list[dict]:
    """Rows as dicts from every sheet, auto-detecting the header row."""
    wb = load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    records: list[dict] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header_idx = None
        for i, row in enumerate(rows[:25]):
            cells = [str(c).strip().lower() for c in row if c is not None]
            hits = sum(any(k in c for k in ("beach", "date", "entero", "station", "site"))
                       for c in cells)
            if hits >= 2:
                header_idx = i
                break
        if header_idx is None:
            continue
        header = [str(c).strip() if c is not None else f"col{j}"
                  for j, c in enumerate(rows[header_idx])]
        for row in rows[header_idx + 1:]:
            if row and any(c is not None for c in row):
                records.append(dict(zip(header, row)))
    return records


def _site_key(rec: dict, name_k, station_k) -> Optional[str]:
    if station_k and rec.get(station_k) not in (None, ""):
        return str(rec[station_k]).strip()
    if name_k and rec.get(name_k):
        return slug(str(rec[name_k]))
    return None


def parse_monitoring_xlsx(xlsx: bytes, detail_url: str = PAGE) -> CollectorResult:
    res = CollectorResult()
    recs = _sheet_records(xlsx)
    if not recs:
        res.warnings.append("VA monitoring workbook: no header row recognized")
        return res
    k = recs[0]
    name_k, station_k = find_key(k, *NAME_KEYS), find_key(k, *STATION_KEYS)
    date_k, result_k = find_key(k, *DATE_KEYS), find_key(k, *RESULT_KEYS)
    lat_k, lon_k, county_k = find_key(k, *LAT_KEYS), find_key(k, *LON_KEYS), find_key(k, *COUNTY_KEYS)
    if station_k == name_k:
        station_k = None
    if not (date_k and result_k and (name_k or station_k)):
        res.warnings.append(f"VA monitoring: unmapped columns {list(k)}")
        return res

    seen: set[str] = set()
    for rec in recs:
        key = _site_key(rec, name_k, station_k)
        d = parse_date(rec.get(date_k))
        if not key or not d:
            continue
        if key not in seen:
            seen.add(key)
            lat = _float(rec.get(lat_k)) if lat_k else None
            lon = _float(rec.get(lon_k)) if lon_k else None
            if lon is not None and lon > 0:  # some sheets drop the minus sign
                lon = -lon
            res.sites.append(Site(
                source_site_id=key, name=str(rec.get(name_k) or key).strip(), state="VA",
                county=str(rec[county_k]).strip() if county_k and rec.get(county_k) else None,
                latitude=lat, longitude=lon, water_type="marine", detail_url=detail_url,
                raw={kk: v for kk, v in rec.items() if kk not in (date_k, result_k)}))
        value, qual = parse_result(rec.get(result_k))
        res.samples.append(Sample(source_site_id=key, sample_date=d,
                                  indicator=Indicator.ENTEROCOCCUS, value=value,
                                  qualifier=qual, units="CFU/100mL", raw=rec))
    return res


def parse_advisory_xlsx(xlsx: bytes) -> CollectorResult:
    res = CollectorResult()
    recs = _sheet_records(xlsx)
    if not recs:
        return res
    k = recs[0]
    name_k, station_k = find_key(k, *NAME_KEYS), find_key(k, *STATION_KEYS)
    if station_k == name_k:
        station_k = None
    start_k, end_k = find_key(k, *ADV_START), find_key(k, *ADV_END)
    type_k, reason_k = find_key(k, *ADV_TYPE), find_key(k, *ADV_REASON)
    for rec in recs:
        key = _site_key(rec, name_k, station_k)
        if not key:
            continue
        type_text = str(rec.get(type_k) or "advisory") if type_k else "advisory"
        kind = "closure" if "clos" in type_text.lower() else "advisory"
        res.advisories.append(Advisory(
            source_site_id=key, kind=kind,
            start_date=parse_date(rec.get(start_k)) if start_k else None,
            end_date=parse_date(rec.get(end_k)) if end_k else None,
            reason=str(rec.get(reason_k)) if reason_k and rec.get(reason_k) else None,
            source_text=type_text, raw=rec))
    return res


def _float(v) -> Optional[float]:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None
