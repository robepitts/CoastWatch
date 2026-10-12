"""Generic, configuration-driven collectors for state and county sources.

Most health departments publish beach results in one of a few shapes:

    table : an HTML table on a web page, or a CSV / Excel download
    json  : a JSON feed (a list of records somewhere inside the document)

Both collectors map each record to our model with a `columns` block in
config/sources.json, so adding a county is usually configuration only:

    "aa_county": {
      "collector": "table",
      "state": "MD",
      "agency": "Anne Arundel County Department of Health",
      "label": "Anne Arundel County Health",
      "url": "https://example.gov/beach-results",
      "table_match": "Enterococci",        # pick the table whose header contains this
      "columns": {
        "name": "Beach",                   # required: beach or station name
        "id": "Station ID",                # optional: stable id (else derived from name)
        "date": "Sample Date",
        "result": "Enterococci",
        "status": "Status",                # optional: published status wording
        "lat": "Latitude", "lon": "Longitude",
        "county": "County", "url": "Link"
      },
      "forward_fill": ["Beach"],           # cells left blank under a merged heading
      "indicator": "enterococcus",         # or "e_coli"
      "priority": 90
    }

Column names are matched loosely (case, spaces and punctuation ignored, partial
matches allowed), so small header changes on the source page don't break it.
Sites without coordinates can be placed with config/coordinate_overrides.csv.
"""
from __future__ import annotations

import csv
import io
import json
import logging
from typing import Any, Optional

from bs4 import BeautifulSoup

from ..models import (CollectorResult, Indicator, Sample, Site, StatusReport,
                      normalize_status)
from .base import Collector, clean_text, find_key, nice_name, parse_date, parse_result, slug

log = logging.getLogger(__name__)

FIELDS = ("id", "name", "date", "result", "status", "lat", "lon", "county", "water_body", "url",
          "units")


# ----------------------------------------------------------------- shared mapping
def rows_to_result(rows: list[dict], columns: dict, state: str, *,
                   indicator: str = "enterococcus", units: Optional[str] = None,
                   detail_url: Optional[str] = None, water_type: str = "marine",
                   forward_fill: Optional[list[str]] = None) -> CollectorResult:
    """Turn a list of row dicts into sites, samples and statuses using `columns`."""
    res = CollectorResult()
    if not rows:
        res.warnings.append("no rows found")
        return res
    keys = _resolve_columns(rows, columns)
    if not keys.get("name") and not keys.get("id"):
        res.warnings.append(f"name/id column not found; columns are: {list(rows[0])}")
        return res
    ind = Indicator.E_COLI if indicator.replace(".", "").replace(" ", "_").lower() in (
        "e_coli", "ecoli") else Indicator.ENTEROCOCCUS

    fill_keys = [find_key(rows[0], c) for c in (forward_fill or [])]
    last: dict[str, Any] = {}
    sites: dict[str, Site] = {}
    latest_status: dict[str, tuple[str, str]] = {}   # sid -> (date, text)
    for row in rows:
        row = dict(row)
        for k in fill_keys:
            if k and clean_text(row.get(k)) is None and k in last:
                row[k] = last[k]
            elif k:
                last[k] = row.get(k)
        get = lambda f: row.get(keys[f]) if keys.get(f) else None   # noqa: E731
        name = clean_text(get("name"))
        sid = clean_text(get("id")) or (slug(name) if name else None)
        if not sid:
            continue
        if sid not in sites:
            lat, lon = _float(get("lat")), _float(get("lon"))
            if lon is not None and lon > 0 and state in ("VA", "MD", "DE", "NJ", "NY"):
                lon = -lon   # western hemisphere; some sheets drop the sign
            url = clean_text(get("url"))
            sites[sid] = Site(
                source_site_id=sid, name=nice_name(name or sid), state=state,
                latitude=lat, longitude=lon, county=nice_name(clean_text(get("county")) or "") or None,
                water_body=clean_text(get("water_body")), water_type=water_type,
                detail_url=url if url and url.startswith("http") else detail_url,
                raw={k: v for k, v in row.items() if k}, )
        d = parse_date(get("date"))
        raw_result = get("result")
        if d and clean_text(raw_result) is not None:
            value, qual = parse_result(raw_result)
            if value is not None or qual:
                u = clean_text(get("units")) or units
                res.samples.append(Sample(
                    source_site_id=sid, sample_date=d, indicator=ind, value=value, qualifier=qual,
                    units=("MPN/100mL" if u and "mpn" in u.lower() else "CFU/100mL"),
                    source_sample_id=f"{d}:{raw_result}", raw=row))
        st = clean_text(get("status"))
        if st:
            prev = latest_status.get(sid)
            if prev is None or (d or "") >= prev[0]:
                latest_status[sid] = (d or "", st)
    res.sites = list(sites.values())
    res.statuses = [StatusReport(sid, normalize_status(t), t) for sid, (_, t) in latest_status.items()]
    return res


def _resolve_columns(rows: list[dict], columns: dict) -> dict[str, Optional[str]]:
    sample = rows[0]
    out = {}
    for f in FIELDS:
        want = columns.get(f)
        if not want:
            continue
        cands = want if isinstance(want, list) else [want]
        out[f] = find_key(sample, *cands)
    return out


def _float(v) -> Optional[float]:
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


# ----------------------------------------------------------------- table collector
def html_tables(html: str) -> list[list[dict]]:
    """Every <table> on the page as a list of row dicts keyed by its header row."""
    soup = BeautifulSoup(html, "lxml")
    out = []
    for t in soup.find_all("table"):
        trs = t.find_all("tr")
        if len(trs) < 2:
            continue
        header = [c.get_text(" ", strip=True) for c in trs[0].find_all(["th", "td"])]
        if not any(header):
            continue
        rows = []
        for tr in trs[1:]:
            cells = [c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])]
            if any(cells):
                rows.append(dict(zip(header, cells + [""] * (len(header) - len(cells)))))
        out.append(rows)
    return out


def csv_rows(text: str) -> list[dict]:
    return [r for r in csv.DictReader(io.StringIO(text)) if any((v or "").strip() for v in r.values())]


def xlsx_rows(data: bytes, sheet: Optional[str] = None, header_hint: Optional[str] = None) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = [wb[sheet]] if sheet else wb.worksheets
    for ws in sheets:
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows[:30]):
            cells = [str(c).strip().lower() for c in row if c is not None]
            if len(cells) >= 2 and (not header_hint or any(header_hint.lower() in c for c in cells)):
                header = [str(c).strip() if c is not None else f"col{j}" for j, c in enumerate(row)]
                return [dict(zip(header, r)) for r in rows[i + 1:] if r and any(c is not None for c in r)]
    return []


def pick_table(tables: list[list[dict]], match: Optional[str]) -> list[dict]:
    if not tables:
        return []
    if match:
        m = match.lower()
        for t in tables:
            if t and any(m in k.lower() for k in t[0]):
                return t
        return []
    return max(tables, key=len)


class TableCollector(Collector):
    """HTML table, CSV or Excel file at a URL."""

    def __init__(self, source_id: str, state: str, agency: str, url: str, columns: dict,
                 homepage_url: str = "", fmt: Optional[str] = None, table_match: Optional[str] = None,
                 sheet: Optional[str] = None, forward_fill: Optional[list[str]] = None,
                 indicator: str = "enterococcus", units: Optional[str] = None,
                 water_type: str = "marine", **kw):
        super().__init__(**kw)
        self.source_id, self.state, self.agency = source_id, state, agency
        self.url, self.columns, self.homepage_url = url, columns, homepage_url or url
        self.fmt, self.table_match, self.sheet = fmt, table_match, sheet
        self.forward_fill, self.indicator, self.units, self.water_type = forward_fill, indicator, units, water_type

    def collect(self) -> CollectorResult:
        resp = self.get(self.url)
        rows = parse_table_payload(resp.body, resp.content_type, self.url, fmt=self.fmt,
                                   table_match=self.table_match, sheet=self.sheet)
        log.info("%s: %d rows; columns: %s", self.source_id, len(rows),
                 ", ".join(list(rows[0])[:25]) if rows else "-")
        res = rows_to_result(rows, self.columns, self.state, indicator=self.indicator,
                             units=self.units, detail_url=self.homepage_url,
                             water_type=self.water_type, forward_fill=self.forward_fill)
        res.warnings = [f"{self.source_id}: {w}" for w in res.warnings]
        return res


def parse_table_payload(body: bytes, content_type: str, url: str, *, fmt: Optional[str] = None,
                        table_match: Optional[str] = None, sheet: Optional[str] = None) -> list[dict]:
    ct, u = (content_type or "").lower(), url.lower().split("?")[0]
    fmt = fmt or ("xlsx" if u.endswith((".xlsx", ".xlsm")) or "spreadsheetml" in ct
                  else "csv" if u.endswith(".csv") or "text/csv" in ct
                  else "html")
    if fmt == "xlsx":
        return xlsx_rows(body, sheet=sheet, header_hint=table_match)
    text = body.decode("utf-8", errors="replace")
    if fmt == "csv":
        return csv_rows(text)
    return pick_table(html_tables(text), table_match)


# ----------------------------------------------------------------- JSON collector
def dig(obj: Any, path: Optional[str]) -> Any:
    """Follow a dotted path like 'data.items' (numbers index into lists)."""
    if not path:
        return obj
    for part in path.split("."):
        if isinstance(obj, list) and part.isdigit():
            obj = obj[int(part)]
        elif isinstance(obj, dict):
            obj = obj.get(part)
        else:
            return None
    return obj


def flatten(rec: dict, prefix: str = "") -> dict:
    """{'loc': {'lat': 1}} -> {'loc.lat': 1} so nested fields can be mapped."""
    out = {}
    for k, v in rec.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(flatten(v, key + "."))
        else:
            out[key] = v
    return out


class JSONCollector(Collector):
    def __init__(self, source_id: str, state: str, agency: str, url: str, columns: dict,
                 records: Optional[str] = None, homepage_url: str = "", params: Optional[dict] = None,
                 indicator: str = "enterococcus", units: Optional[str] = None,
                 water_type: str = "marine", **kw):
        super().__init__(**kw)
        self.source_id, self.state, self.agency = source_id, state, agency
        self.url, self.columns, self.records = url, columns, records
        self.homepage_url, self.params = homepage_url or url, params
        self.indicator, self.units, self.water_type = indicator, units, water_type

    def collect(self) -> CollectorResult:
        data = self.get(self.url, params=self.params).json()
        recs = dig(data, self.records)
        rows = [flatten(r) for r in recs if isinstance(r, dict)] if isinstance(recs, list) else []
        log.info("%s: %d records; fields: %s", self.source_id, len(rows),
                 ", ".join(list(rows[0])[:25]) if rows else "-")
        res = rows_to_result(rows, self.columns, self.state, indicator=self.indicator,
                             units=self.units, detail_url=self.homepage_url, water_type=self.water_type)
        if not isinstance(recs, list):
            res.warnings.append(f"records path {self.records!r} did not lead to a list")
        res.warnings = [f"{self.source_id}: {w}" for w in res.warnings]
        return res
