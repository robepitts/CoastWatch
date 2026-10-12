"""`python -m coastwatch probe` -- inspect what each live source returns.

Run this once from a machine with normal internet access and paste the
output back into the conversation. It reports raw structure (field names,
headers, record counts) without writing to the database, which is what's
needed to pin down field mappings that couldn't be verified remotely.
"""
from __future__ import annotations

import io
import json

from .collectors import CONFIG, build
from .collectors.delaware import parse_embedded_coords, parse_unit_ids
from .collectors.new_jersey import extract_embedded_records
from .http import PoliteClient


def _short(v, n=60):
    s = str(v)
    return s if len(s) <= n else s[:n] + "..."


def probe(only: list[str] | None = None) -> dict:
    cfg_all = json.loads(CONFIG.read_text())
    client = PoliteClient(min_interval_s=1.0)
    report: dict = {}
    for source_id, cfg in cfg_all.items():
        if source_id.startswith("_") or not isinstance(cfg, dict):
            continue
        if only and source_id not in only:
            continue
        col = build(source_id, cfg, client, store=None)
        r: dict = {}
        try:
            kind = cfg["collector"]
            if kind == "arcgis":
                layers = col.resolve_layers(col.item_id) if col.item_id else []
                r["layers"] = []
                for u in layers:
                    meta = client.get(u, params={"f": "json"}).json()
                    feats = col.query_layer(u, page=5)[:1]
                    r["layers"].append({
                        "url": u, "name": meta.get("name"),
                        "fields": [f["name"] for f in meta.get("fields", [])],
                        "example": {k: _short(v) for k, v in
                                    (feats[0].get("properties") or {}).items()} if feats else None,
                    })
                if not layers:
                    r["note"] = "no layers found; try layer_name_filter=null in config"
            elif kind == "virginia":
                from openpyxl import load_workbook
                from .collectors.virginia import PAGE, find_workbook_links
                links = find_workbook_links(client.get(PAGE).text)
                r["links"] = links
                for kind_ in ("monitoring", "advisory"):
                    if links[kind_]:
                        url = links[kind_][max(links[kind_])]
                        wb = load_workbook(io.BytesIO(client.get(url).body), read_only=True,
                                           data_only=True)
                        r[kind_] = {ws.title: [[_short(c, 30) for c in row]
                                               for row in ws.iter_rows(max_row=6, values_only=True)]
                                    for ws in wb.worksheets}
            elif kind == "delaware":
                html = client.get(col.homepage_url).text
                r["discovered_unit_ids"] = sorted(parse_unit_ids(html))
                r["embedded_coords"] = parse_embedded_coords(html)
                r["script_srcs"] = [s for s in __import__("re").findall(
                    r'<script[^>]+src="([^"]+)"', html)][:20]
            elif kind == "new_jersey":
                html = client.get(col.homepage_url).text
                recs = extract_embedded_records(html)
                r["embedded_records"] = len(recs)
                r["example"] = recs[0] if recs else None
                r["has_next_data"] = "__NEXT_DATA__" in html
                r["has_rsc_payload"] = "self.__next_f" in html
            elif kind == "nyc":
                rows = client.get("https://data.cityofnewyork.us/resource/2xir-kwzz.json",
                                  params={"$order": "sample_date DESC", "$limit": 3}).json()
                r["latest_rows"] = rows
            r["ok"] = True
        except Exception as e:
            r["ok"] = False
            r["error"] = f"{type(e).__name__}: {e}"
        report[source_id] = r
    return report
